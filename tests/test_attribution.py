import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
from datafoldit import db, web
from datafoldit.excel_io import export_report_workbook


class AttributionTests(unittest.TestCase):
    def test_all_ledgers_preserve_creator_and_track_editor(self):
        with tempfile.TemporaryDirectory() as tmp:
            conn = db.connect(Path(tmp) / "test.db")
            db.init_db(conn)
            cases = [
                ("bank_transactions", db.add_bank_transaction, db.update_bank_transaction, web.render_bank, {"date": "2026-09-01", "type": "Deposit", "amount": 100}),
                ("expenses", db.add_expense, db.update_expense, web.render_expenses, {"date": "2026-09-01", "amount": 100}),
                ("payroll_entries", db.add_payroll_entry, db.update_payroll_entry, web.render_payroll, {"month": "2026-09", "gross": 100, "first_name": "Sample"}),
                ("invoices", db.add_invoice, db.update_invoice, web.render_invoices, {"date": "2026-09-01", "invoice_number": "ATTR-1", "amount": 100}),
            ]
            for table, create, update, render, payload in cases:
                token = db.ACTOR.set("creator@example.com")
                try:
                    uid = create(conn, {**payload, "created_by": "forged@example.com", "updated_by": "forged@example.com"})
                finally:
                    db.ACTOR.reset(token)
                token = db.ACTOR.set("editor@example.com")
                try:
                    update(conn, uid, payload)
                finally:
                    db.ACTOR.reset(token)
                record = conn.execute(f"SELECT * FROM {table} WHERE id=?", (uid,)).fetchone()
                self.assertEqual(db.attribution_values(record), ["creator@example.com", "editor@example.com"])
                page = render(conn)
                for value in ("Updated by", "editor@example.com"):
                    self.assertIn(value, page)
                self.assertNotIn("Uploaded/Created by", page)
            token = db.ACTOR.set("status-editor@example.com")
            try:
                db.update_invoice_status(conn, 1, "Received")
            finally:
                db.ACTOR.reset(token)
            self.assertEqual(conn.execute("SELECT updated_by FROM invoices").fetchone()[0], "status-editor@example.com")
            # Replay only historical actors actually captured in audit records.
            conn.execute("UPDATE invoices SET created_by=NULL,updated_by=NULL")
            conn.execute("DELETE FROM settings WHERE key='attribution_backfilled_v1'")
            conn.commit()
            db.init_db(conn)
            row = conn.execute("SELECT * FROM invoices").fetchone()
            self.assertEqual(db.attribution_values(row), ["creator@example.com", "status-editor@example.com"])
            path = export_report_workbook(conn, Path(tmp) / "report.xlsx")
            book = load_workbook(path)
            for name in ("Bank Transactions", "Expenses", "Payroll", "Invoices"):
                sheet = book[name]
                self.assertEqual(sheet.cell(1, sheet.max_column).value, "Updated by")
                self.assertEqual(sheet.cell(2, sheet.max_column).value, "status-editor@example.com" if name == "Invoices" else "editor@example.com")
            book.close()
            conn.close()

    def test_unknown_creator_is_not_invented_on_later_update(self):
        with tempfile.TemporaryDirectory() as tmp:
            conn = db.connect(Path(tmp) / "test.db")
            db.init_db(conn)
            uid = db.add_invoice(conn, {"amount": 10})
            token = db.ACTOR.set("editor@example.com")
            try:
                db.update_invoice_status(conn, uid, "Received")
            finally:
                db.ACTOR.reset(token)
            self.assertEqual(db.attribution_values(conn.execute("SELECT * FROM invoices").fetchone()), ["Not recorded", "editor@example.com"])
            conn.close()
